# /// script
# dependencies = ["openpyxl"]
# ///
# Offline guard for the Dizzle Dynasty sheet parse. No network.
#   uv run .claude/skills/dizzle-dynasty/test_sheet.py
#   uv run .claude/skills/dizzle-dynasty/test_sheet.py --record   (re-cut fixtures from live)
#
# Runs sheet.py against a recorded real workbook with urlopen stubbed, diffs its output
# against a recorded snapshot, then removes each guard in turn and asserts the removal
# does serve the wrong board — so the guards are never "simplified" away.
import difflib
import io
import pathlib
import runpy
import sys

try:                                    # noqa: E402 -- must precede any openpyxl use
    import openpyxl                     # noqa: F401
except ModuleNotFoundError:
    sys.exit("run this with `uv run`, not `python3` -- the PEP 723 block above declares\n"
             "openpyxl and uv resolves it. This is NOT a broken test: three separate agents\n"
             "have reported it as one after invoking it with python3.\n"
             "  uv run .claude/skills/dizzle-dynasty/test_sheet.py")
import urllib.request
from contextlib import redirect_stdout

HERE = pathlib.Path(__file__).parent
SCRIPT = HERE / 'sheet.py'
FIX = HERE / 'fixtures'
BOOK = FIX / 'dizzle-dynasty.xlsx'
SNAP = FIX / 'dizzle-dynasty.expected.txt'

PULLS = [[], ['Dynasty Ranks, Points', '20'], ['Pick Values', '60']]

# The documented ways this recipe silently serves the wrong board, as source mutations.
TAKE_FIRST_MATCH = ("assert len(hits) == 1", "assert len(hits) >= 1")
ALLOW_HIDDEN = ("assert ws.sheet_state == 'visible'",
                "assert ws.sheet_state in ('visible', 'hidden')")
DROP_PICK_FORMAT = ("    if pick and isinstance(v, (int, float)): return f'{v:.2f}'\n", '')
NAIVE_PICK_FORMAT = ("if pick and isinstance(v, (int, float)): return f'{v:.2f}'",
                     "if pick and isinstance(v, (int, float)): return str(v)")


def run_script(args, source=None):
    """Run sheet.py with urlopen stubbed to serve the fixture workbook. Returns stdout."""
    book = BOOK.read_bytes()
    real, argv = urllib.request.urlopen, sys.argv
    urllib.request.urlopen = lambda url, timeout=None: io.BytesIO(book)
    sys.argv = [str(SCRIPT), *args]
    buf = io.StringIO()
    try:
        with redirect_stdout(buf):
            try:
                if source is None:
                    runpy.run_path(str(SCRIPT), run_name='__main__')
                else:
                    exec(compile(source, str(SCRIPT), 'exec'), {'__name__': '__main__'})
            except SystemExit:
                pass
    finally:
        urllib.request.urlopen, sys.argv = real, argv
    return buf.getvalue()


def render():
    return ''.join(f'$ sheet.py {" ".join(a)}\n{run_script(a)}\n' for a in PULLS)


def mutate(find, replace):
    src = SCRIPT.read_text()
    out = src.replace(find, replace, 1)
    assert out != src, f'mutation no longer applies, source moved: {find}'
    return out


def test_ambiguous_tab_name_is_refused_not_silently_resolved_to_9cat():
    try:
        run_script(['Dynasty Ranks'])
    except AssertionError as e:
        assert 'be more specific' in str(e) and '9Cat' in str(e), e
    else:
        raise AssertionError('"Dynasty Ranks" matched one tab — fixture lost a board')
    got = run_script(['Dynasty Ranks'], mutate(*TAKE_FIRST_MATCH))
    assert 'Dynasty Ranks, 9Cat  [visible]' in got and 'FORMAT   9CAT' in got, \
        f'taking hits[0] no longer serves 9Cat — guard is dead\n{got[:400]}'


def test_hidden_tab_is_refused_not_served_as_a_current_board():
    hidden = next(l.split(None, 1)[1] for l in run_script([]).splitlines()
                  if l.startswith('hidden'))
    try:
        run_script([hidden])
    except AssertionError as e:
        assert 'stale archive' in str(e), e
    else:
        raise AssertionError(f'"{hidden}" passed the visible check — guard is dead')
    got = run_script([hidden], mutate(*ALLOW_HIDDEN))
    assert '[hidden]' in got and '\t#\tPlayer\t' in got, \
        f'hidden archive no longer serves its old column layout\n{got[:400]}'


def test_float_pick_labels_keep_their_two_decimal_slot_form():
    got = run_script(['Pick Values', '60'])
    assert '\n1.10\t' in got and '\n2.30\t' in got, f'slot labels lost\n{got[:400]}'
    dropped = run_script(['Pick Values', '60'], mutate(*DROP_PICK_FORMAT))
    assert '\n1.1\t' in dropped and '\n2.3\t' in dropped, \
        f'dropping the pick case no longer degrades the label — Pick is not float-typed ' \
        f'in the fixture\n{dropped[:400]}'
    naive = run_script(['Pick Values', '60'], mutate(*NAIVE_PICK_FORMAT))
    assert '\n1.1\t' in naive and '\n2.30000000000004\t' in naive, \
        f'a naive str() no longer shows the float noise\n{naive[:400]}'


def test_parses_the_recorded_pulls():
    got, want = render(), SNAP.read_bytes().decode()
    assert got == want, '\n' + '\n'.join(difflib.unified_diff(
        want.splitlines(), got.splitlines(), 'recorded', 'now', lineterm=''))


def record():
    """Re-cut the fixture from the live sheet: a real download with rows trimmed."""
    import openpyxl
    raw = urllib.request.urlopen(
        'https://docs.google.com/spreadsheets/d/'
        '1EmReTa5KUcFFMCy8Fq-NpQG3WU0pMY7XNW54G3EPbEM/export?format=xlsx', timeout=90).read()
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    for ws in wb.worksheets:
        keep = (60 if ws.title.strip() == 'Pick Values'
                else 20 if 'Dynasty Ranks' in ws.title else 6)
        if ws.max_row > keep + 1:
            ws.delete_rows(keep + 2, ws.max_row - keep - 1)
    FIX.mkdir(exist_ok=True)
    wb.save(BOOK)
    SNAP.write_bytes(render().encode())
    print(f'{[w.title for w in wb.worksheets]}\n{BOOK} ({BOOK.stat().st_size} B)\n{SNAP}')


if __name__ == '__main__':
    if '--record' in sys.argv[1:]:
        record()
        sys.exit()
    tests = [v for k, v in sorted(globals().items()) if k.startswith('test_')]
    for t in tests:
        t()
        print(f'ok  {t.__name__}')
    print(f'{len(tests)} passed')
